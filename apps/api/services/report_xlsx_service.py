"""A live report as a spreadsheet a CA can sort, filter and add up (accounting-16).

THE SECOND RENDERER of the `ReportDocument` the PDF is drawn from
(`services/report_pdf_service.py`), so the two cannot disagree with each other
and neither can disagree with the screen the figures were copied from.

THE RULES, all of which `apps/web/lib/export/xlsx.ts` already holds for the
browser's own exports and are restated here because this is the second
implementation of them and a spreadsheet is exactly where a reader adds a
column up:

    A MONEY CELL IS A NUMBER. A figure written as text is skipped by `=SUM()`,
    so a CA who selects the amount column of an exported ledger and asks for a
    total gets 0 with nothing to say why. The value is RUPEES as a number —
    `Decimal(paise) / 100`, exact for every integer, never a float — and the
    format string tells Excel how to DISPLAY it.

    INDIAN GROUPING IN THE FORMAT STRING (decision D6). Excel has no
    locale-free way to say it, so it is the three-condition form the Indian
    Excel community uses: above a crore, above a lakh, and below. The string is
    the same one the browser exports use; `tests/test_report_export_*` asserts
    the two are equal so one cannot drift.

    A FIGURE NOBODY HOLDS IS AN EMPTY CELL, NOT A ZERO. `None` writes nothing.

    THE HEADER ROW FREEZES AND THE COLUMNS ARE WIDE ENOUGH. A ledger whose
    particulars are clipped to eight characters is not a ledger.

A BALANCE KEEPS ITS SIGN. The PDF prints a ledger balance as `1,000.00 Dr`; a
spreadsheet cannot add that up, so the cell is the signed number, debit-positive,
and the header says so.

THE DATA SHEETS CARRY DATA ONLY. The practice, the client, the period and the
caveats go on a sheet of their own ("Details") so that every data sheet has its
header on row 1 and can be pivoted without deleting a title block first; the
print header of each data sheet still names the client and the report.
"""
from __future__ import annotations

import io
import re
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core.ist_clock import ist_today
from domain.firm.identity import gstin_of
from domain.reporting.export_builders import sheet_name as _safe_sheet_name
from domain.reporting.export_document import (
    BALANCE, EMPHASIS, INT, MONEY, SECTION, ReportDocument, cell_text,
)
from services import pdf_style

#: Excel's number format for Indian grouping with two decimals. The same string
#: as `INR_FORMAT` in apps/web/lib/export/xlsx.ts, pinned equal by a test.
INR_FORMAT = '[>=10000000]##\\,##\\,##\\,##0.00;[>=100000]##\\,##\\,##0.00;##,##0.00'

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_HUNDRED = Decimal(100)
_MIN_WIDTH, _MAX_WIDTH = 10, 60
_WIDTH_SAMPLE = 500          # rows read to size a column; the rest cannot widen it much


def rupees(paise: int) -> Decimal:
    """Integer paise as rupees, exactly. A Decimal, never a float."""
    return Decimal(int(paise)) / _HUNDRED


def _hex(token: str) -> str:
    return token.lstrip("#")


def _label(col) -> str:
    return f"{col.label} (Dr +, Cr -)" if col.kind == BALANCE else col.label


def _unique_sheet(wb: Workbook, wanted: str) -> str:
    """A sheet name Excel accepts that is not already in the workbook."""
    base = _safe_sheet_name(wanted)
    name, n = base, 2
    existing = {ws.title for ws in wb.worksheets}
    while name in existing:
        suffix = f" ({n})"
        name = base[: 31 - len(suffix)] + suffix
        n += 1
    return name


def build_report_xlsx(doc: ReportDocument, *, holder: dict, firm: dict,
                      generated_on: str | None = None) -> bytes:
    client_name = (holder.get("legal_name") or holder.get("trade_name")
                   or holder.get("client_name") or "")
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.title = f"{doc.title} - {client_name}".strip(" -")
    wb.properties.creator = firm.get("name") or ""

    head_fill = PatternFill("solid", fgColor=_hex(pdf_style.INK))
    head_font = Font(bold=True, color="FFFFFF")
    band_fill = PatternFill("solid", fgColor=_hex(pdf_style.MUTED))
    bold = Font(bold=True)

    for index, tbl in enumerate(doc.tables):
        if len(doc.tables) == 1:
            wanted = doc.sheet_name
        else:
            wanted = tbl.heading or f"{doc.sheet_name} {index + 1}"
        ws = wb.create_sheet(_unique_sheet(wb, wanted))
        ws.append([_label(c) for c in tbl.columns])
        for cell in ws[1]:
            cell.font, cell.fill = head_font, head_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)

        widths = [len(_label(c)) for c in tbl.columns]
        for r, row in enumerate(tbl.rows, start=2):
            values = []
            for i, (col, value) in enumerate(zip(tbl.columns, row.cells)):
                if value is None:
                    values.append(None)
                elif col.kind in (MONEY, BALANCE):
                    values.append(rupees(value))
                elif col.kind == INT:
                    values.append(int(value))
                else:
                    values.append(value)
                if r - 2 < _WIDTH_SAMPLE:
                    widths[i] = max(widths[i], len(cell_text(col.kind, value)))
            ws.append(values)
            for i, col in enumerate(tbl.columns, start=1):
                cell = ws.cell(row=r, column=i)
                if col.kind in (MONEY, BALANCE):
                    cell.number_format = INR_FORMAT
                if row.style in (EMPHASIS, SECTION):
                    cell.font = bold
                    cell.fill = band_fill
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = max(
                _MIN_WIDTH, min(_MAX_WIDTH, w + 2))
        ws.freeze_panes = "A2"
        ws.oddHeader.center.text = f"{client_name} - {doc.title}".strip(" -")
        ws.page_setup.orientation = "landscape" if doc.landscape else "portrait"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = "1:1"

    # ── The sheet that says what this is ─────────────────────────────────────
    info = wb.create_sheet(_unique_sheet(wb, "Details"))
    firm_gstin = gstin_of(firm)
    lines: list[tuple[str, str]] = [
        ("Prepared by", firm.get("name") or ""),
        ("Practice GSTIN", firm_gstin or ""),
        ("Client", client_name),
        ("Client GSTIN", holder.get("gstin") or ""),
        ("Client PAN", holder.get("pan") or ""),
        ("Report", doc.title + (f" - {doc.subject}" if doc.subject else "")),
        ("Period", doc.period),
        ("Currency", "INR; amounts are rupees, held to the paisa"),
        ("Generated on", f"{generated_on or ist_today().isoformat()} (IST)"),
    ]
    lines.extend(doc.meta)
    for key, value in lines:
        if value:
            info.append([key, value])
    info.append([])
    for note in doc.notes:
        info.append(["Note", note])
    info.column_dimensions["A"].width = 18
    info.column_dimensions["B"].width = 100
    for row in info.iter_rows(min_col=1, max_col=1):
        for cell in row:
            if cell.value:
                cell.font = bold
    for row in info.iter_rows(min_col=2, max_col=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


_FILENAME_SAFE = re.compile(r"[^A-Za-z0-9._-]+")


def filename_for(stem: str, ext: str) -> str:
    return f"{_FILENAME_SAFE.sub('-', stem).strip('-') or 'report'}.{ext}"
